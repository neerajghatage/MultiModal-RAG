"""
Multi-format file loader for Documentation documentation.

Supports: PDF, HTML/HTM, Images (PNG/JPG/GIF/BMP/PSD/AI),
Excel (XLS/XLSX), XML, CSS, JS, and MadCap WebHelp files.
"""

from pathlib import Path
import logging
import hashlib
import base64
import os
import io
import csv
import json
import warnings
import re
from typing import List, Dict, Any, Optional
from datetime import datetime

from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from PIL import Image

# PyMuPDF (fitz) for high-fidelity PDF extraction; fall back to pypdf
try:
    import fitz as pymupdf  # PyMuPDF
    _HAS_PYMUPDF = True
except ImportError:
    _HAS_PYMUPDF = False

from pypdf import PdfReader  # fallback

logger = logging.getLogger(__name__)

MAX_FILE_SIZE_BYTES = 200 * 1024 * 1024
SKIP_DIRS = {".git", "__pycache__", "node_modules", ".svn", ".hg"}


class MultiFileLoader:
    """
    Load documents from Documentation:
    - PDF (text + embedded images)
    - HTML/HTM (text + image references)
    - PNG/JPG/JPEG/GIF/BMP/PSD/AI (images for VLM summarization)
    - XLSX/XLS (Excel spreadsheets)
    - XML (configuration/structured data)
    - CSS/JS (stylesheets and scripts)
    - MCWEBHELP (MadCap Flare help config)
    """

    SUPPORTED_EXTENSIONS = {
        "pdf": [".pdf"],
        "html": [".html", ".htm", ".mcwebhelp"],
        "image": [".png", ".jpg", ".jpeg", ".gif", ".bmp", ".psd", ".ai"],
        "excel": [".xlsx", ".xls"],
        "structured_text": [".xml"],
        "code": [".css", ".js"],
    }

    def __init__(self):
        self.loaded_files: List[Path] = []
        self.skipped_files: List[Dict[str, str]] = []
        self.error_files: List[Dict[str, str]] = []
        self.file_stats = {
            "pdf": 0, "html": 0, "image": 0,
            "excel": 0, "code": 0, "structured_text": 0, "total": 0,
        }
        logger.info("MultiFileLoader initialized")

    # ── Helpers ───────────────────────────────────────────────

    @staticmethod
    def _generate_file_id(filepath: Path) -> str:
        return f"file_{hashlib.sha256(str(filepath).encode()).hexdigest()[:12]}"

    def _get_file_type(self, filepath: Path) -> str:
        ext = filepath.suffix.lower()
        for ftype, exts in self.SUPPORTED_EXTENSIONS.items():
            if ext in exts:
                return ftype
        return "unknown"

    @staticmethod
    def _is_safe_to_load(filepath: Path) -> Optional[str]:
        try:
            size = filepath.stat().st_size
        except OSError as e:
            return f"Cannot stat file: {e}"
        if size == 0:
            return "Zero-byte file"
        if size > MAX_FILE_SIZE_BYTES:
            return f"File too large ({size / (1024**2):.0f} MB)"
        return None

    @staticmethod
    def _read_text_with_fallback(filepath: Path) -> str:
        for enc in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
            try:
                with open(filepath, "r", encoding=enc) as f:
                    return f.read()
            except (UnicodeDecodeError, UnicodeError):
                continue
        with open(filepath, "rb") as f:
            return f.read().decode("utf-8", errors="replace")

    def _chunk_text(self, file_id: str, content: str, max_chunk_size: int = 2000) -> List[Dict[str, Any]]:
        chunks = []
        if len(content) > max_chunk_size:
            paragraphs = content.split("\n\n")
            current_chunk: List[str] = []
            current_len = 0
            chunk_idx = 0
            for para in paragraphs:
                if para.strip():
                    if current_len + len(para) > max_chunk_size and current_chunk:
                        chunks.append({
                            "element_id": f"{file_id}_text_{chunk_idx}",
                            "content": "\n\n".join(current_chunk),
                            "type": "text",
                        })
                        chunk_idx += 1
                        current_chunk = []
                        current_len = 0
                    current_chunk.append(para)
                    current_len += len(para)
            if current_chunk:
                chunks.append({
                    "element_id": f"{file_id}_text_{chunk_idx}",
                    "content": "\n\n".join(current_chunk),
                    "type": "text",
                })
        else:
            chunks.append({"element_id": f"{file_id}_text_0", "content": content, "type": "text"})
        return chunks

    # ── Individual loaders ────────────────────────────────────

    def load_pdf(self, filepath: Path) -> Dict[str, Any]:
        logger.info(f"Loading PDF: {filepath.name}")
        skip = self._is_safe_to_load(filepath)
        if skip:
            return {"error": skip, "filename": filepath.name, "filepath": str(filepath)}

        file_id = self._generate_file_id(filepath)

        # ── Try PyMuPDF first (higher-fidelity extraction) ────
        if _HAS_PYMUPDF:
            try:
                return self._load_pdf_pymupdf(filepath, file_id)
            except Exception as e:
                logger.warning(f"PyMuPDF failed for {filepath.name}, falling back to pypdf: {e}")

        # ── Fallback to pypdf ─────────────────────────────────
        try:
            return self._load_pdf_pypdf(filepath, file_id)
        except Exception as e:
            logger.error(f"Error loading PDF {filepath}: {e}")
            return {"error": str(e), "filename": filepath.name, "filepath": str(filepath)}

    def _load_pdf_pymupdf(self, filepath: Path, file_id: str) -> Dict[str, Any]:
        """High-fidelity PDF extraction using PyMuPDF (fitz)."""
        doc = pymupdf.open(str(filepath))
        text_chunks = []
        images = []

        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text("text")
            if text.strip():
                text_chunks.append({
                    "element_id": f"{file_id}_page_{page_num}",
                    "content": text,
                    "type": "text",
                    "page": page_num + 1,
                })

            # Extract embedded images
            for img_idx, img_info in enumerate(page.get_images(full=True)):
                try:
                    xref = img_info[0]
                    pix = pymupdf.Pixmap(doc, xref)
                    if pix.n > 4:  # CMYK → RGB
                        pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
                    b64 = base64.b64encode(pix.tobytes("png")).decode("utf-8")
                    images.append({
                        "element_id": f"{file_id}_img_{page_num}_{img_idx}",
                        "content": b64,
                        "image_base64": b64,
                        "type": "image",
                        "format": "png",
                        "page": page_num + 1,
                    })
                except Exception:
                    pass

        page_count = len(doc)
        doc.close()

        if not text_chunks:
            text_chunks.append({
                "element_id": f"{file_id}_page_0",
                "content": f"[PDF with {page_count} page(s) — no extractable text]",
                "type": "text", "page": 0,
            })

        self.file_stats["pdf"] += 1
        self.file_stats["total"] += 1
        return {
            "file_id": file_id, "filename": filepath.name, "filepath": str(filepath),
            "type": "pdf", "text_chunks": text_chunks, "tables": [], "images": images,
            "page_count": page_count, "element_count": len(text_chunks) + len(images),
            "loaded_at": datetime.now().isoformat(),
        }

    def _load_pdf_pypdf(self, filepath: Path, file_id: str) -> Dict[str, Any]:
        """Fallback PDF extraction using pypdf."""
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            reader = PdfReader(str(filepath))

        text_chunks = []
        for page_num, page in enumerate(reader.pages):
            try:
                text = page.extract_text() or ""
            except Exception:
                text = ""
            if text.strip():
                text_chunks.append({
                    "element_id": f"{file_id}_page_{page_num}",
                    "content": text, "type": "text", "page": page_num + 1,
                })

        if not text_chunks:
            text_chunks.append({
                "element_id": f"{file_id}_page_0",
                "content": f"[PDF with {len(reader.pages)} page(s) — no extractable text]",
                "type": "text", "page": 0,
            })

        self.file_stats["pdf"] += 1
        self.file_stats["total"] += 1
        return {
            "file_id": file_id, "filename": filepath.name, "filepath": str(filepath),
            "type": "pdf", "text_chunks": text_chunks, "tables": [], "images": [],
            "page_count": len(reader.pages), "element_count": len(text_chunks),
            "loaded_at": datetime.now().isoformat(),
        }

    def load_docx(self, filepath: Path) -> Dict[str, Any]:
        logger.info(f"Loading DOCX: {filepath.name}")
        skip = self._is_safe_to_load(filepath)
        if skip:
            return {"error": skip, "filename": filepath.name, "filepath": str(filepath)}
        try:
            doc = DocxDocument(filepath)
            file_id = self._generate_file_id(filepath)
            text_chunks = []
            images = []

            for idx, para in enumerate(doc.paragraphs):
                if para.text.strip():
                    text_chunks.append({
                        "element_id": f"{file_id}_text_{idx}",
                        "content": para.text,
                        "type": "text",
                        "style": para.style.name if para.style else "Normal",
                    })

            for idx, rel in enumerate(doc.part.rels.values()):
                if "image" in rel.target_ref:
                    try:
                        image_data = rel.target_part.blob
                        b64 = base64.b64encode(image_data).decode("utf-8")
                        images.append({
                            "element_id": f"{file_id}_image_{idx}",
                            "content": b64,
                            "image_base64": b64,
                            "type": "image",
                            "format": rel.target_ref.split(".")[-1],
                        })
                    except Exception:
                        pass

            self.file_stats["docx"] += 1
            self.file_stats["total"] += 1
            return {
                "file_id": file_id, "filename": filepath.name, "filepath": str(filepath),
                "type": "docx", "text_chunks": text_chunks, "images": images,
                "element_count": len(text_chunks) + len(images),
                "loaded_at": datetime.now().isoformat(),
            }
        except Exception as e:
            logger.error(f"Error loading DOCX {filepath}: {e}")
            return {"error": str(e), "filename": filepath.name, "filepath": str(filepath)}

    def load_doc_legacy(self, filepath: Path) -> Dict[str, Any]:
        logger.info(f"Loading legacy DOC: {filepath.name}")
        skip = self._is_safe_to_load(filepath)
        if skip:
            return {"error": skip, "filename": filepath.name, "filepath": str(filepath)}

        file_id = self._generate_file_id(filepath)
        try:
            import textract
            content = textract.process(str(filepath)).decode("utf-8", errors="replace")
        except ImportError:
            logger.warning(f"textract not installed; extracting strings from .doc: {filepath.name}")
            try:
                with open(filepath, "rb") as f:
                    raw = f.read()
                content = "\n".join(
                    m.group().decode("ascii", errors="replace")
                    for m in re.finditer(rb"[\x20-\x7E]{4,}", raw)
                )
            except Exception as e:
                return {"error": str(e), "filename": filepath.name, "filepath": str(filepath)}
        except Exception as e:
            return {"error": str(e), "filename": filepath.name, "filepath": str(filepath)}

        if not content.strip():
            return {"error": "No readable text from .doc", "filename": filepath.name, "filepath": str(filepath)}

        chunks = self._chunk_text(file_id, content)
        self.file_stats["doc_legacy"] += 1
        self.file_stats["total"] += 1
        return {
            "file_id": file_id, "filename": filepath.name, "filepath": str(filepath),
            "type": "doc_legacy", "text_chunks": chunks,
            "element_count": len(chunks), "loaded_at": datetime.now().isoformat(),
        }

    def load_html(self, filepath: Path) -> Dict[str, Any]:
        logger.info(f"Loading HTML: {filepath.name}")
        skip = self._is_safe_to_load(filepath)
        if skip:
            return {"error": skip, "filename": filepath.name, "filepath": str(filepath)}
        try:
            raw = self._read_text_with_fallback(filepath)
            file_id = self._generate_file_id(filepath)

            # Use heading-aware chunker for structured HTML parsing
            from chunker import HeadingAwareChunker
            chunker = HeadingAwareChunker()
            chunks, referenced_images = chunker.chunk_html(
                filepath=str(filepath), raw_html=raw, file_id=file_id
            )

            if not chunks:
                chunks.append({
                    "element_id": f"{file_id}_text_0",
                    "content": "[HTML file with no extractable text]",
                    "type": "text",
                    "chunk_type": "parent",
                    "parent_chunk_id": None,
                    "hierarchy_path": "",
                    "heading_level": 0,
                    "section_title": filepath.stem,
                })

            self.file_stats["html"] += 1
            self.file_stats["total"] += 1
            return {
                "file_id": file_id, "filename": filepath.name, "filepath": str(filepath),
                "type": "html", "text_chunks": chunks,
                "referenced_images": referenced_images,
                "element_count": len(chunks), "loaded_at": datetime.now().isoformat(),
            }
        except Exception as e:
            logger.error(f"Error loading HTML {filepath}: {e}")
            return {"error": str(e), "filename": filepath.name, "filepath": str(filepath)}

    def load_image(self, filepath: Path) -> Dict[str, Any]:
        logger.info(f"Loading Image: {filepath.name}")
        skip = self._is_safe_to_load(filepath)
        if skip:
            return {"error": skip, "filename": filepath.name, "filepath": str(filepath)}
        try:
            with Image.open(filepath) as img:
                if img.mode != "RGB":
                    img = img.convert("RGB")
                buf = io.BytesIO()
                img.save(buf, format="PNG")
                file_id = self._generate_file_id(filepath)
                b64 = base64.b64encode(buf.getvalue()).decode("utf-8")
                self.file_stats["image"] += 1
                self.file_stats["total"] += 1
                return {
                    "file_id": file_id, "filename": filepath.name, "filepath": str(filepath),
                    "type": "image",
                    "images": [{
                        "element_id": f"{file_id}_image_0",
                        "content": f"Image: {filepath.name} ({img.width}x{img.height})",
                        "image_data": b64,
                        "image_base64": b64,
                        "type": "image",
                        "metadata": {"dimensions": f"{img.width}x{img.height}", "format": img.format},
                    }],
                    "dimensions": f"{img.width}x{img.height}", "format": img.format,
                    "element_count": 1, "loaded_at": datetime.now().isoformat(),
                }
        except Exception as e:
            logger.error(f"Error loading image {filepath}: {e}")
            return {"error": str(e), "filename": filepath.name, "filepath": str(filepath)}

    def load_text(self, filepath: Path) -> Dict[str, Any]:
        logger.info(f"Loading Text: {filepath.name}")
        skip = self._is_safe_to_load(filepath)
        if skip:
            return {"error": skip, "filename": filepath.name, "filepath": str(filepath)}
        try:
            content = self._read_text_with_fallback(filepath)
            file_id = self._generate_file_id(filepath)
            chunks = self._chunk_text(file_id, content)
            self.file_stats["text"] += 1
            self.file_stats["total"] += 1
            return {
                "file_id": file_id, "filename": filepath.name, "filepath": str(filepath),
                "type": "text", "text_chunks": chunks,
                "element_count": len(chunks), "loaded_at": datetime.now().isoformat(),
            }
        except Exception as e:
            logger.error(f"Error loading text {filepath}: {e}")
            return {"error": str(e), "filename": filepath.name, "filepath": str(filepath)}

    def load_code(self, filepath: Path) -> Dict[str, Any]:
        logger.info(f"Loading Code: {filepath.name}")
        skip = self._is_safe_to_load(filepath)
        if skip:
            return {"error": skip, "filename": filepath.name, "filepath": str(filepath)}
        try:
            content = self._read_text_with_fallback(filepath)
            file_id = self._generate_file_id(filepath)
            self.file_stats["code"] += 1
            self.file_stats["total"] += 1
            return {
                "file_id": file_id, "filename": filepath.name, "filepath": str(filepath),
                "type": "code", "language": filepath.suffix[1:],
                "text_chunks": [{
                    "element_id": f"{file_id}_code_0",
                    "content": content,
                    "type": "code",
                    "metadata": {"language": filepath.suffix[1:]},
                }],
                "lines": len(content.split("\n")), "element_count": 1,
                "loaded_at": datetime.now().isoformat(),
            }
        except Exception as e:
            logger.error(f"Error loading code {filepath}: {e}")
            return {"error": str(e), "filename": filepath.name, "filepath": str(filepath)}

    def load_csv(self, filepath: Path) -> Dict[str, Any]:
        logger.info(f"Loading CSV: {filepath.name}")
        skip = self._is_safe_to_load(filepath)
        if skip:
            return {"error": skip, "filename": filepath.name, "filepath": str(filepath)}
        try:
            content = self._read_text_with_fallback(filepath)
            file_id = self._generate_file_id(filepath)
            reader = csv.reader(io.StringIO(content))
            rows = list(reader)
            if not rows:
                return {"error": "Empty CSV", "filename": filepath.name, "filepath": str(filepath)}

            headers = rows[0]
            text_chunks = [{"element_id": f"{file_id}_header", "content": f"CSV Headers: {', '.join(headers)}", "type": "table_header"}]
            data_rows = rows[1:]
            batch_size = 50
            for bi in range(0, len(data_rows), batch_size):
                batch = data_rows[bi : bi + batch_size]
                lines = []
                for row in batch:
                    row_dict = {h: v for h, v in zip(headers, row)} if headers else {}
                    lines.append(", ".join(f"{k}: {v}" for k, v in row_dict.items()) if row_dict else ", ".join(row))
                text_chunks.append({
                    "element_id": f"{file_id}_rows_{bi}",
                    "content": "\n".join(lines),
                    "type": "table",
                    "row_range": f"{bi + 1}-{bi + len(batch)}",
                })

            self.file_stats["csv"] += 1
            self.file_stats["total"] += 1
            return {
                "file_id": file_id, "filename": filepath.name, "filepath": str(filepath),
                "type": "csv", "text_chunks": text_chunks, "row_count": len(data_rows),
                "column_count": len(headers), "element_count": len(text_chunks),
                "loaded_at": datetime.now().isoformat(),
            }
        except Exception as e:
            logger.error(f"Error loading CSV {filepath}: {e}")
            return {"error": str(e), "filename": filepath.name, "filepath": str(filepath)}

    def load_excel(self, filepath: Path) -> Dict[str, Any]:
        logger.info(f"Loading Excel: {filepath.name}")
        skip = self._is_safe_to_load(filepath)
        if skip:
            return {"error": skip, "filename": filepath.name, "filepath": str(filepath)}

        file_id = self._generate_file_id(filepath)
        text_chunks: List[Dict[str, Any]] = []

        try:
            import openpyxl
            wb = openpyxl.load_workbook(str(filepath), read_only=True, data_only=True)
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows = list(ws.iter_rows(values_only=True))
                if not rows:
                    continue
                headers = [str(c) if c is not None else "" for c in rows[0]]
                data_rows = rows[1:]
                batch_size = 50
                for bi in range(0, max(len(data_rows), 1), batch_size):
                    batch = data_rows[bi : bi + batch_size]
                    lines = []
                    for row in batch:
                        row_dict = {h: str(v) if v is not None else "" for h, v in zip(headers, row)}
                        lines.append(", ".join(f"{k}: {v}" for k, v in row_dict.items() if v))
                    text_chunks.append({
                        "element_id": f"{file_id}_{sheet_name}_{bi}",
                        "content": f"Sheet: {sheet_name}\nHeaders: {', '.join(headers)}\n" + "\n".join(lines),
                        "type": "table", "sheet": sheet_name,
                    })
            wb.close()
        except ImportError:
            try:
                import pandas as pd
                xls = pd.ExcelFile(str(filepath))
                for sn in xls.sheet_names:
                    df = pd.read_excel(xls, sheet_name=sn)
                    text_chunks.append({
                        "element_id": f"{file_id}_{sn}_0",
                        "content": f"Sheet: {sn}\n{df.to_string(max_rows=200)}",
                        "type": "table", "sheet": sn,
                    })
            except Exception as e2:
                return {"error": str(e2), "filename": filepath.name, "filepath": str(filepath)}
        except Exception as e:
            if filepath.suffix.lower() == ".xls":
                try:
                    import pandas as pd
                    xls = pd.ExcelFile(str(filepath))
                    for sn in xls.sheet_names:
                        df = pd.read_excel(xls, sheet_name=sn)
                        text_chunks.append({
                            "element_id": f"{file_id}_{sn}_0",
                            "content": f"Sheet: {sn}\n{df.to_string(max_rows=200)}",
                            "type": "table", "sheet": sn,
                        })
                except Exception as e2:
                    return {"error": str(e2), "filename": filepath.name, "filepath": str(filepath)}
            else:
                return {"error": str(e), "filename": filepath.name, "filepath": str(filepath)}

        if not text_chunks:
            return {"error": "No data in Excel", "filename": filepath.name, "filepath": str(filepath)}

        self.file_stats["excel"] += 1
        self.file_stats["total"] += 1
        return {
            "file_id": file_id, "filename": filepath.name, "filepath": str(filepath),
            "type": "excel", "text_chunks": text_chunks,
            "element_count": len(text_chunks), "loaded_at": datetime.now().isoformat(),
        }

    def load_structured_text(self, filepath: Path) -> Dict[str, Any]:
        logger.info(f"Loading structured text: {filepath.name}")
        skip = self._is_safe_to_load(filepath)
        if skip:
            return {"error": skip, "filename": filepath.name, "filepath": str(filepath)}
        try:
            content = self._read_text_with_fallback(filepath)
            file_id = self._generate_file_id(filepath)
            chunks = self._chunk_text(file_id, content)
            self.file_stats["structured_text"] += 1
            self.file_stats["total"] += 1
            return {
                "file_id": file_id, "filename": filepath.name, "filepath": str(filepath),
                "type": "structured_text", "sub_type": filepath.suffix.lower().lstrip("."),
                "text_chunks": chunks, "element_count": len(chunks),
                "loaded_at": datetime.now().isoformat(),
            }
        except Exception as e:
            logger.error(f"Error loading structured text {filepath}: {e}")
            return {"error": str(e), "filename": filepath.name, "filepath": str(filepath)}

    # ── Dispatch ──────────────────────────────────────────────

    def load_file(self, filepath: Path) -> Dict[str, Any]:
        ftype = self._get_file_type(filepath)
        loaders = {
            "pdf": self.load_pdf,
            "html": self.load_html, "image": self.load_image, "code": self.load_code,
            "excel": self.load_excel,
            "structured_text": self.load_structured_text,
        }
        loader = loaders.get(ftype)
        if loader:
            return loader(filepath)
        self.skipped_files.append({"filename": filepath.name, "filepath": str(filepath), "reason": f"Unsupported: {filepath.suffix}"})
        return {"error": f"Unsupported file type: {filepath.suffix}", "filename": filepath.name, "filepath": str(filepath)}

    def load_all_documents(self, folder_path: Path, recursive: bool = True) -> List[Dict[str, Any]]:
        warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
        warnings.filterwarnings("ignore", category=UserWarning, module="PIL")

        documents = []
        folder_path = Path(folder_path)

        if not folder_path.exists():
            logger.error(f"Folder does not exist: {folder_path}")
            return documents

        all_files = []
        for root, dirs, files in os.walk(folder_path):
            dirs[:] = [d for d in dirs if d not in SKIP_DIRS]
            if not recursive and Path(root) != folder_path:
                continue
            for f in files:
                all_files.append(Path(root) / f)

        supported = [f for f in all_files if self._get_file_type(f) != "unknown"]
        unsupported = [f for f in all_files if self._get_file_type(f) == "unknown"]
        for f in unsupported:
            self.skipped_files.append({"filename": f.name, "filepath": str(f), "reason": f"Unsupported: {f.suffix}"})

        logger.info(f"Found {len(all_files)} files — {len(supported)} supported, {len(unsupported)} skipped")

        for idx, fp in enumerate(supported):
            if (idx + 1) % 500 == 0:
                logger.info(f"  Progress: {idx + 1}/{len(supported)}")
            doc = self.load_file(fp)
            if "error" not in doc:
                documents.append(doc)
                self.loaded_files.append(fp)
            else:
                self.error_files.append({"filename": fp.name, "filepath": str(fp), "error": doc.get("error", "")})

        self._print_stats()
        return documents

    def _print_stats(self):
        logger.info(f"\n{'='*60}\nLOADING STATISTICS\n{'='*60}")
        for ftype, count in sorted(self.file_stats.items()):
            if ftype != "total" and count > 0:
                logger.info(f"  {ftype:20s}: {count}")
        logger.info(f"  {'TOTAL':20s}: {self.file_stats['total']}")
        if self.skipped_files:
            logger.info(f"  {'SKIPPED':20s}: {len(self.skipped_files)}")
        if self.error_files:
            logger.info(f"  {'ERRORS':20s}: {len(self.error_files)}")
        logger.info("=" * 60)

    def get_loading_report(self) -> Dict[str, Any]:
        return {
            "stats": dict(self.file_stats),
            "loaded_count": len(self.loaded_files),
            "skipped_count": len(self.skipped_files),
            "error_count": len(self.error_files),
            "skipped_files": self.skipped_files[:50],
            "error_files": self.error_files[:50],
        }
